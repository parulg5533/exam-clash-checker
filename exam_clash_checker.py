"""
Exam Clash Checker Algorithm & Tool
------------------------------------
Analyzes student course registration data and exam timetables to detect:
1. Hard Clashes: Multiple exams scheduled for the same student in the exact same time slot.
2. Soft Warnings: Two exams on the same day (morning & evening) for a single student.
3. Multi-slot anomalies: A course appearing in more than one exam slot.
4. Unscheduled courses: Registered courses missing from the timetable.
5. Safe slot recommendations: Identify conflict-free slots for any specific course.
"""

from collections import defaultdict
from dataclasses import dataclass, field
import os
import openpyxl
import pandas as pd


@dataclass
class CourseInfo:
    code: str
    name: str
    enrolled_count: int = 0
    students: set = field(default_factory=set)


@dataclass
class SlotClash:
    date_str: str
    slot_name: str
    course_a: str
    course_b: str
    overlap_count: int
    students: list


@dataclass
class TimetableAuditResult:
    total_slots: int
    total_courses_in_timetable: int
    clashes: list  # list of SlotClash
    duplicate_courses: dict  # course -> list of slots
    unscheduled_courses: list  # list of course codes
    same_day_warnings: list  # list of dicts with student, date, morning_course, evening_course
    students_with_clashes: dict  # student -> list of clashed slots


class ExamClashChecker:
    def __init__(self, reg_data_path: str = None, timetable_path: str = None):
        self.reg_data_path = reg_data_path or "Autumn 2026 Registration Data as on 31-8-2026.xlsx"
        self.timetable_path = timetable_path or "Tentative TimeTable Autumn 2026 IITP v1.xlsx"

        self.student_to_courses = defaultdict(set)
        self.course_to_students = defaultdict(set)
        self.course_info = {}
        self.conflict_graph = defaultdict(dict)  # c1 -> {c2: set(students)}
        
        self.slots = []  # list of dict: {id, date_str, session, full_name, courses}
        self.course_to_slots = defaultdict(list)

        self._load_registration_data()
        self._build_conflict_graph()
        if os.path.exists(self.timetable_path):
            self.load_timetable(self.timetable_path)

    def _load_registration_data(self):
        """Loads registration data and course titles from the Excel workbook."""
        wb = openpyxl.load_workbook(self.reg_data_path, data_only=True)
        
        # 1. Course names from 'Sub Name' if present
        if "Sub Name" in wb.sheetnames:
            ws_sub = wb["Sub Name"]
            for row in list(ws_sub.iter_rows(values_only=True))[1:]:
                if row and row[0]:
                    c_code = str(row[0]).strip()
                    c_name = str(row[1]).strip() if len(row) > 1 and row[1] else c_code
                    self.course_info[c_code] = CourseInfo(code=c_code, name=c_name)

        # 2. Student registrations from 'Reg Data'
        ws_reg = wb["Reg Data"] if "Reg Data" in wb.sheetnames else wb.active
        for row in list(ws_reg.iter_rows(values_only=True))[1:]:
            if not row or not row[0]:
                continue
            roll = str(row[0]).strip()
            # Courses are in Sub1 to Sub8 (columns index 1 to 8)
            courses = [str(c).strip() for c in row[1:9] if c is not None and str(c).strip() != ""]
            
            for c in courses:
                self.student_to_courses[roll].add(c)
                self.course_to_students[c].add(roll)
                if c not in self.course_info:
                    self.course_info[c] = CourseInfo(code=c, name=c)

        # Update enrollment counts
        for c, stus in self.course_to_students.items():
            self.course_info[c].enrolled_count = len(stus)
            self.course_info[c].students = stus

    def _build_conflict_graph(self):
        """Precomputes the student overlap between all pairs of courses."""
        courses = list(self.course_to_students.keys())
        n = len(courses)
        for i in range(n):
            c1 = courses[i]
            stus1 = self.course_to_students[c1]
            for j in range(i + 1, n):
                c2 = courses[j]
                stus2 = self.course_to_students[c2]
                overlap = stus1.intersection(stus2)
                if overlap:
                    self.conflict_graph[c1][c2] = overlap
                    self.conflict_graph[c2][c1] = overlap

    def load_timetable(self, timetable_path: str):
        """Loads and parses the timetable Excel file."""
        self.timetable_path = timetable_path
        self.slots.clear()
        self.course_to_slots.clear()

        wb = openpyxl.load_workbook(timetable_path, data_only=True)
        ws = wb["Timetable"] if "Timetable" in wb.sheetnames else wb.active

        slot_index = 0
        for row in list(ws.iter_rows(values_only=True))[1:]:
            date_val = row[0]
            if not date_val or str(date_val).strip() == "None":
                continue
            
            if hasattr(date_val, "strftime"):
                date_str = date_val.strftime("%Y-%m-%d")
            else:
                date_str = str(date_val).strip().split(" ")[0]

            # Morning slot
            morning_raw = row[1]
            if morning_raw and str(morning_raw).strip() not in ["None", "`", ""]:
                m_courses = self._parse_course_list(str(morning_raw))
                slot_id = f"slot_{slot_index}"
                slot_name = "Morning (10:30 am - 12:30 pm)"
                slot_dict = {
                    "id": slot_id,
                    "date": date_str,
                    "session": "Morning",
                    "slot_name": slot_name,
                    "full_name": f"{date_str} {slot_name}",
                    "courses": m_courses,
                }
                self.slots.append(slot_dict)
                for c in m_courses:
                    self.course_to_slots[c].append(slot_dict)
                slot_index += 1

            # Evening slot
            evening_raw = row[2] if len(row) > 2 else None
            if evening_raw and str(evening_raw).strip() not in ["None", "`", ""]:
                e_courses = self._parse_course_list(str(evening_raw))
                slot_id = f"slot_{slot_index}"
                slot_name = "Evening (3:30 pm - 5:30 pm)"
                slot_dict = {
                    "id": slot_id,
                    "date": date_str,
                    "session": "Evening",
                    "slot_name": slot_name,
                    "full_name": f"{date_str} {slot_name}",
                    "courses": e_courses,
                }
                self.slots.append(slot_dict)
                for c in e_courses:
                    self.course_to_slots[c].append(slot_dict)
                slot_index += 1

    @staticmethod
    def _parse_course_list(cell_text: str) -> list:
        """Parses comma and slash-separated course codes in a timetable cell."""
        courses = []
        for part in cell_text.split(","):
            for subpart in part.split("/"):
                c = subpart.strip()
                if c and c not in ["`", ""]:
                    courses.append(c)
        return courses

    def audit_timetable(self) -> TimetableAuditResult:
        """Performs a comprehensive clash and sanity audit on the loaded timetable."""
        clashes = []
        students_with_clashes = defaultdict(list)

        # 1. Slot-level direct clashes
        for slot in self.slots:
            courses = slot["courses"]
            for i in range(len(courses)):
                c1 = courses[i]
                for j in range(i + 1, len(courses)):
                    c2 = courses[j]
                    if c1 == c2:
                        continue
                    overlap = self.conflict_graph[c1].get(c2, set())
                    if overlap:
                        clash_obj = SlotClash(
                            date_str=slot["date"],
                            slot_name=slot["slot_name"],
                            course_a=c1,
                            course_b=c2,
                            overlap_count=len(overlap),
                            students=sorted(list(overlap)),
                        )
                        clashes.append(clash_obj)
                        for roll in overlap:
                            students_with_clashes[roll].append({
                                "slot": slot["full_name"],
                                "course_a": c1,
                                "course_b": c2
                            })

        # 2. Courses scheduled multiple times
        duplicate_courses = {}
        for c, sls in self.course_to_slots.items():
            if len(sls) > 1:
                duplicate_courses[c] = [s["full_name"] for s in sls]

        # 3. Unscheduled courses
        all_reg_courses = set(self.course_to_students.keys())
        scheduled_courses = set(self.course_to_slots.keys())
        unscheduled = sorted(list(all_reg_courses - scheduled_courses))

        # 4. Same-day exam fatigue (morning + evening)
        day_map = defaultdict(lambda: {"Morning": [], "Evening": []})
        for slot in self.slots:
            day_map[slot["date"]][slot["session"]].extend(slot["courses"])

        same_day_warnings = []
        for date_str, sess in day_map.items():
            m_courses = sess["Morning"]
            e_courses = sess["Evening"]
            if not m_courses or not e_courses:
                continue
            
            # Check students who have at least one course in morning and one in evening
            for mc in m_courses:
                stus_mc = self.course_to_students.get(mc, set())
                if not stus_mc:
                    continue
                for ec in e_courses:
                    stus_ec = self.course_to_students.get(ec, set())
                    common = stus_mc.intersection(stus_ec)
                    if common:
                        same_day_warnings.append({
                            "date": date_str,
                            "morning_course": mc,
                            "evening_course": ec,
                            "overlap_count": len(common),
                            "students": sorted(list(common)),
                        })

        return TimetableAuditResult(
            total_slots=len(self.slots),
            total_courses_in_timetable=len(scheduled_courses),
            clashes=clashes,
            duplicate_courses=duplicate_courses,
            unscheduled_courses=unscheduled,
            same_day_warnings=same_day_warnings,
            students_with_clashes=dict(students_with_clashes),
        )

    def get_course_conflicts(self, course_code: str) -> list:
        """Returns all courses that share enrolled students with the given course."""
        course_code = course_code.strip()
        conflicts = []
        if course_code not in self.course_to_students:
            return conflicts

        for other_course, stus in self.conflict_graph[course_code].items():
            conflicts.append({
                "course_code": other_course,
                "course_name": self.course_info.get(other_course, CourseInfo(code=other_course, name=other_course)).name,
                "overlap_count": len(stus),
                "overlapping_students": sorted(list(stus)),
            })
        
        # Sort descending by overlap count
        conflicts.sort(key=lambda x: x["overlap_count"], reverse=True)
        return conflicts

    def check_course_in_slot(self, course_code: str, slot_index_or_id) -> dict:
        """Checks if a given course clashes with courses in a specific timetable slot."""
        course_code = course_code.strip()
        target_slot = None
        if isinstance(slot_index_or_id, int) and 0 <= slot_index_or_id < len(self.slots):
            target_slot = self.slots[slot_index_or_id]
        else:
            for s in self.slots:
                if s["id"] == str(slot_index_or_id) or s["full_name"] == str(slot_index_or_id):
                    target_slot = s
                    break

        if not target_slot:
            return {"error": f"Slot '{slot_index_or_id}' not found"}

        slot_clashes = []
        all_affected_students = set()

        for c in target_slot["courses"]:
            if c == course_code:
                continue
            overlap = self.conflict_graph[course_code].get(c, set())
            if overlap:
                slot_clashes.append({
                    "clashing_course": c,
                    "course_name": self.course_info.get(c, CourseInfo(code=c, name=c)).name,
                    "overlap_count": len(overlap),
                    "students": sorted(list(overlap)),
                })
                all_affected_students.update(overlap)

        return {
            "course_code": course_code,
            "slot_id": target_slot["id"],
            "slot_name": target_slot["full_name"],
            "has_clash": len(slot_clashes) > 0,
            "clashes": slot_clashes,
            "total_affected_students": len(all_affected_students),
            "affected_students": sorted(list(all_affected_students)),
        }

    def find_safe_slots(self, course_code: str) -> dict:
        """Finds all conflict-free slots in the timetable for a given course."""
        course_code = course_code.strip()
        safe_slots = []
        clashing_slots = []

        for idx, slot in enumerate(self.slots):
            res = self.check_course_in_slot(course_code, idx)
            slot_data = {
                "slot_id": slot["id"],
                "date": slot["date"],
                "session": slot["session"],
                "full_name": slot["full_name"],
                "current_course_count": len(slot["courses"]),
            }
            if not res["has_clash"]:
                safe_slots.append(slot_data)
            else:
                slot_data["clash_summary"] = [
                    f"{c['clashing_course']} ({c['overlap_count']} stds)" for c in res["clashes"]
                ]
                slot_data["total_affected"] = res["total_affected_students"]
                clashing_slots.append(slot_data)

        return {
            "course_code": course_code,
            "course_name": self.course_info.get(course_code, CourseInfo(code=course_code, name=course_code)).name,
            "total_slots": len(self.slots),
            "safe_slots_count": len(safe_slots),
            "safe_slots": safe_slots,
            "clashing_slots_count": len(clashing_slots),
            "clashing_slots": clashing_slots,
        }

    def get_student_schedule(self, roll_number: str) -> dict:
        """Retrieves the complete exam timetable and clash status for a specific student."""
        roll_number = roll_number.strip().upper()
        enrolled = sorted(list(self.student_to_courses.get(roll_number, set())))
        if not enrolled:
            return {"error": f"Student '{roll_number}' not found in registration data."}

        schedule = []
        slot_counts = defaultdict(list)
        day_sessions = defaultdict(lambda: {"Morning": [], "Evening": []})

        for c in enrolled:
            sls = self.course_to_slots.get(c, [])
            c_name = self.course_info.get(c, CourseInfo(code=c, name=c)).name
            if not sls:
                schedule.append({
                    "course_code": c,
                    "course_name": c_name,
                    "status": "Unscheduled",
                    "slot": None,
                })
            else:
                for s in sls:
                    schedule.append({
                        "course_code": c,
                        "course_name": c_name,
                        "status": "Scheduled",
                        "slot": s["full_name"],
                        "date": s["date"],
                        "session": s["session"],
                    })
                    slot_counts[s["full_name"]].append(c)
                    day_sessions[s["date"]][s["session"]].append(c)

        # Identify any clashes for this student
        student_clashes = []
        for slot_name, courses in slot_counts.items():
            if len(courses) > 1:
                student_clashes.append({
                    "slot": slot_name,
                    "courses": courses,
                    "message": f"Double-booked in {slot_name}: {', '.join(courses)}"
                })

        # Identify same day exams
        same_day_exams = []
        for d, sess in day_sessions.items():
            if sess["Morning"] and sess["Evening"]:
                same_day_exams.append({
                    "date": d,
                    "morning": sess["Morning"],
                    "evening": sess["Evening"]
                })

        return {
            "roll_number": roll_number,
            "total_enrolled": len(enrolled),
            "enrolled_courses": enrolled,
            "has_clashes": len(student_clashes) > 0,
            "clashes": student_clashes,
            "same_day_exams": same_day_exams,
            "schedule": schedule,
        }

    def export_audit_to_excel(self, output_filepath: str = "Exam_Clash_Audit_Report.xlsx"):
        """Exports a complete, multi-tab audit report to an Excel workbook."""
        audit = self.audit_timetable()

        with pd.ExcelWriter(output_filepath, engine="openpyxl") as writer:
            # 1. Summary sheet
            summary_data = [
                {"Metric": "Total Registered Students", "Value": len(self.student_to_courses)},
                {"Metric": "Total Registered Courses", "Value": len(self.course_to_students)},
                {"Metric": "Total Exam Slots", "Value": audit.total_slots},
                {"Metric": "Total Courses Scheduled", "Value": audit.total_courses_in_timetable},
                {"Metric": "Total Hard Slot Clashes", "Value": len(audit.clashes)},
                {"Metric": "Students Double-Booked in Slot", "Value": len(audit.students_with_clashes)},
                {"Metric": "Courses in Multiple Slots", "Value": len(audit.duplicate_courses)},
                {"Metric": "Unscheduled Courses", "Value": len(audit.unscheduled_courses)},
                {"Metric": "Same-Day Back-to-Back Warnings", "Value": len(audit.same_day_warnings)},
            ]
            pd.DataFrame(summary_data).to_excel(writer, sheet_name="Audit Summary", index=False)

            # 2. Hard Clashes
            if audit.clashes:
                clash_rows = []
                for cl in audit.clashes:
                    clash_rows.append({
                        "Date": cl.date_str,
                        "Slot": cl.slot_name,
                        "Course A": cl.course_a,
                        "Course A Name": self.course_info.get(cl.course_a, CourseInfo(code=cl.course_a, name=cl.course_a)).name,
                        "Course B": cl.course_b,
                        "Course B Name": self.course_info.get(cl.course_b, CourseInfo(code=cl.course_b, name=cl.course_b)).name,
                        "Overlap Count": cl.overlap_count,
                        "Affected Student Rolls": ", ".join(cl.students),
                    })
                pd.DataFrame(clash_rows).to_excel(writer, sheet_name="Direct Slot Clashes", index=False)
            else:
                pd.DataFrame([{"Status": "No direct slot clashes detected!"}]).to_excel(writer, sheet_name="Direct Slot Clashes", index=False)

            # 3. Double Booked Students
            if audit.students_with_clashes:
                stud_clash_rows = []
                for roll, cl_list in audit.students_with_clashes.items():
                    for cl in cl_list:
                        stud_clash_rows.append({
                            "Roll Number": roll,
                            "Slot": cl["slot"],
                            "Course A": cl["course_a"],
                            "Course B": cl["course_b"],
                        })
                pd.DataFrame(stud_clash_rows).to_excel(writer, sheet_name="Affected Students", index=False)

            # 4. Duplicate Courses
            if audit.duplicate_courses:
                dup_rows = [{"Course Code": c, "Scheduled Slots": " | ".join(sls)} for c, sls in audit.duplicate_courses.items()]
                pd.DataFrame(dup_rows).to_excel(writer, sheet_name="Duplicate Entries", index=False)

            # 5. Same Day Exam Warnings
            if audit.same_day_warnings:
                sd_rows = []
                for w in audit.same_day_warnings:
                    sd_rows.append({
                        "Date": w["date"],
                        "Morning Course": w["morning_course"],
                        "Evening Course": w["evening_course"],
                        "Overlap Count": w["overlap_count"],
                        "Affected Rolls (Sample)": ", ".join(w["students"][:15]) + ("..." if len(w["students"]) > 15 else "")
                    })
                pd.DataFrame(sd_rows).to_excel(writer, sheet_name="Same-Day Fatigue", index=False)

        return output_filepath
